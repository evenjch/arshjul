"""Antall sett i BOM-en (B10) følger årshjulet – ellers regner Excel feil «Å bestille»."""
import datetime as dt

import openpyxl
import pytest

import bom
import config
import frister
from hjelp import lag_bom

G = "/api/gjennomforinger"
NESTE = frister.skolear_for(dt.date.today()) + 1  # ikke startet
KOPIER = config.BOM_DIR / "Sikkerhetskopi"


def b10(path):
    return openpyxl.load_workbook(path).worksheets[0]["B10"].value


def prosjekt(client, skolear=NESTE, sett_i_fil=12):
    """61 studenter i grupper på 4 + 2 ekstra = 18 sett, med en BOM som sier noe annet."""
    g = client.post(G, json={"navn": "Test", "skolear": skolear, "start": 44, "studenter": 61, "gruppe": 4,
                             "ekstra": 2}).get_json()
    fil = lag_bom(rader=[("Arduino", "Elfa", "A1", 1, 0, 250.0)], sett=sett_i_fil)
    client.patch(f"{G}/{g['id']}", json={"bom": fil.name})
    return g, fil


def test_write_sett(tmp_path):
    fil = lag_bom(rader=[("Arduino", "Elfa", "A1", 2, 1, 10.0)], sett=10)
    res = bom.write_sett(fil, 18, KOPIER, forventet=10)
    assert (res["skrevet"], res["for"], res["fil"]) == (True, 10, 18)
    ws = openpyxl.load_workbook(fil).worksheets[0]
    assert ws["B10"].value == 18 and ws["I14"].value.startswith("=")  # formlene står
    assert (KOPIER / res["backup"]).exists() and b10(KOPIER / res["backup"]) == 10
    assert bom.read_bom(fil)["rader"][0]["lager"] == 1  # ingenting annet er endret

    assert bom.write_sett(fil, 18, KOPIER)["skrevet"] is False  # allerede riktig: ingen kopi
    res = bom.write_sett(fil, 20, KOPIER, forventet=10)  # noen har endret B10 siden
    assert (res["skrevet"], res.get("konflikt"), b10(fil)) == (False, True, 18)
    assert bom.write_sett(fil, 20, KOPIER)["skrevet"] is True  # uten forventet: overskriv

    (fil.parent / f"~${fil.name}").write_text("")  # åpen i Excel
    with pytest.raises(bom.BomError, match="åpen i Excel"):
        bom.write_sett(fil, 21, KOPIER)


def test_endret_antall_skrives_i_bom(client):
    g, fil = prosjekt(client, sett_i_fil=18)
    r = client.patch(f"{G}/{g['id']}", json={"gruppe": 3}).get_json()  # 21 + 2 = 23
    assert r["sett"] == 23 and r["bomSett"]["skrevet"] is True
    assert b10(fil) == 23 and bom.read_bom(fil)["sett"] == 23
    felt = [(e["felt"], e["fra"], e["til"]) for e in client.get("/api/admin/logg").get_json()]
    assert ("antall sett i BOM-en", "18", "23") in felt

    r = client.patch(f"{G}/{g['id']}", json={"notat": "x"}).get_json()  # antall sett uendret
    assert r["bomSett"] is None


def test_endret_for_hand_i_excel_overskrives_ikke(client):
    g, fil = prosjekt(client, sett_i_fil=12)  # årshjulet: 18, BOM-en: 12
    st = client.get("/api/state").get_json()["prosjekter"][0]
    assert (st["sett"], st["bomStatus"]["settFil"]) == (18, 12)  # avviket vises
    r = client.patch(f"{G}/{g['id']}", json={"ekstra": 3}).get_json()
    assert r["bomSett"]["konflikt"] and "endret i Excel" in r["bomSett"]["grunn"] and b10(fil) == 12

    r = client.post(f"{G}/{g['id']}/bom/sett").get_json()  # knappen: ingeniøren bestemmer
    assert (r["skrevet"], r["fil"], b10(fil)) == (True, 19, 19)


def test_apen_i_excel_endrer_bare_arshjulet(client):
    g, fil = prosjekt(client, sett_i_fil=18)
    (fil.parent / f"~${fil.name}").write_text("")
    r = client.patch(f"{G}/{g['id']}", json={"gruppe": 3})
    assert r.status_code == 200 and r.get_json()["sett"] == 23
    assert "åpen i Excel" in r.get_json()["bomSett"]["grunn"] and b10(fil) == 18


def test_startet_prosjekt_endrer_ikke_bom(client):
    g, fil = prosjekt(client, skolear=NESTE - 2, sett_i_fil=18)
    r = client.patch(f"{G}/{g['id']}", json={"gruppe": 3}).get_json()
    assert r["bomSett"]["skrevet"] is False and "startet" in r["bomSett"]["grunn"] and b10(fil) == 18


def test_siden_viser_avvik_med_knapp(client):
    from test_side import kjor_siden
    prosjekt(client, sett_i_fil=12)
    res = kjor_siden(client.get("/api/state").get_json())
    assert res["errors"] == []
    assert "BOM-en har 12 sett" in res["tbody"] and "bomavvik" in res["tbody"] and 'data-act="bomsett"' in res["tbody"] and "Skriv 18 til BOM-en" in res["tbody"]


def test_bare_neste_gjennomforing_folger_bom_en(client):
    """Flere skoleår deler BOM-en: B10 følger den neste som ikke har startet og ikke er telt."""
    gammel, fil = prosjekt(client, skolear=NESTE - 2, sett_i_fil=18)  # ferdig for lenge siden
    ny = client.post(G, json={"navn": "Test", "skolear": NESTE, "start": 44, "studenter": 40,
                              "bom": fil.name}).get_json()
    senere = client.post(G, json={"navn": "Test", "skolear": NESTE + 1, "start": 44, "studenter": 50,
                                  "bom": fil.name}).get_json()
    aktiv = {p["id"]: p["bomSettAktiv"] for p in client.get("/api/state").get_json()["prosjekter"]}
    assert aktiv == {gammel["id"]: False, ny["id"]: True, senere["id"]: False}  # ingen falsk alarm på de andre

    r = client.patch(f"{G}/{senere['id']}", json={"studenter": 55}).get_json()["bomSett"]
    assert r["skrevet"] is False and r["feil"] is False and "følger" in r["grunn"] and b10(fil) == 18
    assert client.post(f"{G}/{ny['id']}/bom/sett").get_json()["fil"] == 40  # B10 hadde fjorårets 18: knappen
    assert client.patch(f"{G}/{ny['id']}", json={"studenter": 41}).get_json()["bomSett"]["skrevet"] is True
    assert b10(fil) == 41

    client.post(f"{G}/{ny['id']}/telling", json={"linjer": [{"key": "elfa|a1", "telt": 1}]})  # telt: ferdig
    aktiv = {p["id"]: p["bomSettAktiv"] for p in client.get("/api/state").get_json()["prosjekter"]}
    assert aktiv[senere["id"]] is True and aktiv[ny["id"]] is False


def test_ferdig_gjennomforing_viser_ikke_avvik(client):
    from test_side import kjor_siden
    prosjekt(client, skolear=NESTE - 2, sett_i_fil=12)  # startet: B10 bryr seg ikke om den lenger
    res = kjor_siden(client.get("/api/state").get_json())
    assert res["errors"] == [] and "bomavvik" not in res["tbody"]
