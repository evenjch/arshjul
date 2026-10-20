"""The set count in the BOM (B10) follows the app – otherwise Excel computes «Å bestille» wrong."""
import datetime as dt

import openpyxl
import pytest

import bom
import config
import deadlines
from helpers import make_bom

G = "/api/runs"
NEXT_SY = deadlines.school_year_for(dt.date.today()) + 1  # not started
BACKUPS = config.BOM_DIR / "Sikkerhetskopi"


def b10(path):
    return openpyxl.load_workbook(path).worksheets[0]["B10"].value


def project(client, school_year=NEXT_SY, sets_in_file=12):
    """61 students in groups of 4 + 2 spare = 18 sets, with a BOM that says something else."""
    g = client.post(G, json={"name": "Test", "school_year": school_year, "start": 44, "students": 61, "group_size": 4,
                             "spare": 2}).get_json()
    file = make_bom(rows=[("Arduino", "Elfa", "A1", 1, 0, 250.0)], sets=sets_in_file)
    client.patch(f"{G}/{g['id']}", json={"bom": file.name})
    return g, file


def test_write_sets(tmp_path):
    file = make_bom(rows=[("Arduino", "Elfa", "A1", 2, 1, 10.0)], sets=10)
    res = bom.write_sets(file, 18, BACKUPS, expected=10)
    assert (res["written"], res["for"], res["file"]) == (True, 10, 18)
    ws = openpyxl.load_workbook(file).worksheets[0]
    assert ws["B10"].value == 18 and ws["I14"].value.startswith("=")  # the formulas are intact
    assert (BACKUPS / res["backup"]).exists() and b10(BACKUPS / res["backup"]) == 10
    assert bom.read_bom(file)["rows"][0]["in_stock"] == 1  # nothing else has changed

    assert bom.write_sets(file, 18, BACKUPS)["written"] is False  # already right: no backup
    res = bom.write_sets(file, 20, BACKUPS, expected=10)  # someone has changed B10 since
    assert (res["written"], res.get("conflict"), b10(file)) == (False, True, 18)
    assert bom.write_sets(file, 20, BACKUPS)["written"] is True  # without expected: overwrite

    (file.parent / f"~${file.name}").write_text("")  # open in Excel
    with pytest.raises(bom.BomError, match="åpen i Excel"):
        bom.write_sets(file, 21, BACKUPS)


def test_changed_count_is_written_to_the_bom(client):
    g, file = project(client, sets_in_file=18)
    r = client.patch(f"{G}/{g['id']}", json={"group_size": 3}).get_json()  # 21 + 2 = 23
    assert r["sets"] == 23 and r["bomSets"]["written"] is True
    assert b10(file) == 23 and bom.read_bom(file)["sets"] == 23
    fields = [(e["field"], e["from_value"], e["to_value"]) for e in client.get("/api/admin/log").get_json()]
    assert ("antall sett i BOM-en", "18", "23") in fields

    r = client.patch(f"{G}/{g['id']}", json={"note": "x"}).get_json()  # set count unchanged
    assert r["bomSets"] is None


def test_hand_edit_in_excel_is_not_overwritten(client):
    g, file = project(client, sets_in_file=12)  # the app: 18, the BOM: 12
    st = client.get("/api/state").get_json()["projects"][0]
    assert (st["sets"], st["bomStatus"]["setsInFile"]) == (18, 12)  # avviket vises
    r = client.patch(f"{G}/{g['id']}", json={"spare": 3}).get_json()
    assert r["bomSets"]["conflict"] and "endret i Excel" in r["bomSets"]["reason"] and b10(file) == 12

    r = client.post(f"{G}/{g['id']}/bom/sets").get_json()  # the button: the engineer decides
    assert (r["written"], r["file"], b10(file)) == (True, 19, 19)


def test_open_in_excel_changes_only_the_app(client):
    g, file = project(client, sets_in_file=18)
    (file.parent / f"~${file.name}").write_text("")
    r = client.patch(f"{G}/{g['id']}", json={"group_size": 3})
    assert r.status_code == 200 and r.get_json()["sets"] == 23
    assert "åpen i Excel" in r.get_json()["bomSets"]["reason"] and b10(file) == 18


def test_started_project_doesnt_change_bom(client):
    g, file = project(client, school_year=NEXT_SY - 2, sets_in_file=18)
    r = client.patch(f"{G}/{g['id']}", json={"group_size": 3}).get_json()
    assert r["bomSets"]["written"] is False and "startet" in r["bomSets"]["reason"] and b10(file) == 18


def test_page_shows_mismatch_with_button(client):
    from test_page import run_page
    project(client, sets_in_file=12)
    res = run_page(client.get("/api/state").get_json())
    assert res["errors"] == []
    assert "BOM-en har 12 sett" in res["tbody"] and "bom-mismatch" in res["tbody"] and 'data-act="bomsets"' in res["tbody"] and "Skriv 18 til BOM-en" in res["tbody"]


def test_only_the_next_run_follows_the_bom(client):
    """Several school years share the BOM: B10 follows the next one that hasn't started and hasn't been counted."""
    old_count, file = project(client, school_year=NEXT_SY - 2, sets_in_file=18)  # ferdig for lenge siden
    new = client.post(G, json={"name": "Test", "school_year": NEXT_SY, "start": 44, "students": 40,
                              "bom": file.name}).get_json()
    later = client.post(G, json={"name": "Test", "school_year": NEXT_SY + 1, "start": 44, "students": 50,
                                  "bom": file.name}).get_json()
    active = {p["id"]: p["bomSetsActive"] for p in client.get("/api/state").get_json()["projects"]}
    assert active == {old_count["id"]: False, new["id"]: True, later["id"]: False}  # no false alarm on the others

    r = client.patch(f"{G}/{later['id']}", json={"students": 55}).get_json()["bomSets"]
    assert r["written"] is False and r["error"] is False and "følger" in r["reason"] and b10(file) == 18
    assert client.post(f"{G}/{new['id']}/bom/sets").get_json()["file"] == 40  # B10 had last year's 18: the button
    assert client.patch(f"{G}/{new['id']}", json={"students": 41}).get_json()["bomSets"]["written"] is True
    assert b10(file) == 41

    client.post(f"{G}/{new['id']}/count", json={"lines": [{"key": "elfa|a1", "counted": 1}]})  # telt: ferdig
    active = {p["id"]: p["bomSetsActive"] for p in client.get("/api/state").get_json()["projects"]}
    assert active[later["id"]] is True and active[new["id"]] is False


def test_finished_run_shows_no_mismatch(client):
    from test_page import run_page
    project(client, school_year=NEXT_SY - 2, sets_in_file=12)  # started: B10 no longer cares about it
    res = run_page(client.get("/api/state").get_json())
    assert res["errors"] == [] and "bom-mismatch" not in res["tbody"]
