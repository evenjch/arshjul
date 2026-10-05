"""Felles hjelpere for testene."""
import datetime as dt

import openpyxl

import bom
import config


def lag_bom(navn="Test – BOM.xlsx", rader=(), sett=10):
    """Lag en BOM fra malen i test-datamappen.

    rader: (beskrivelse, leverandør, varenummer, per sett, på lager, pris)
    """
    target = config.BOM_DIR / navn
    bom.make_template(config.TEMPLATE_BOM, target, prosjekt="Testprosjekt", dato=dt.date(2026, 8, 1),
                      frist=dt.date(2026, 6, 15), sett=sett)
    if rader:
        skriv_rader(target, rader)
    return target


def skriv_rader(target, rader):
    wb = openpyxl.load_workbook(target)
    ws = wb.worksheets[0]
    for r in range(bom.FIRST_ROW, bom.LAST_ROW + 1):
        for c in "CEFGHJ":
            ws[f"{c}{r}"] = None
    for i, (besk, lev, varenr, pr_sett, lager, pris) in enumerate(rader):
        r = bom.FIRST_ROW + i
        ws[f"C{r}"], ws[f"E{r}"], ws[f"F{r}"] = besk, lev, varenr
        ws[f"G{r}"], ws[f"H{r}"], ws[f"J{r}"] = pr_sett, lager, pris
    wb.save(target)
