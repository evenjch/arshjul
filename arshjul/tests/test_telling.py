import openpyxl
import pytest

import bom
import config
import telling
from hjelp import lag_bom

G = "/api/gjennomforinger"


def rad(besk, lev, pr_sett, lager, bestill, pris=10.0, varenr=""):
    return {"beskrivelse": besk, "leverandor": lev, "varenummer": varenr, "prSett": pr_sett, "lager": lager,
            "bestill": bestill, "pris": pris, "total": bestill * pris}


# ---------- logikken ----------

def test_grunnlag_bruker_registrert_bestilling_eller_antar_bom():
    rader = [rad("Arduino", "Elfa", 1, 4, 6, varenr="A1"), rad("Ledning", "Kjell", 2, 0, 20, varenr="L1")]
    # Elfa er registrert (og det ble bestilt 8), Kjell ikke
    g = telling.grunnlag(rader, ({"elfa|a1": 8.0}, {"elfa"}), {"kjell|l1": "forbruk"}, sett=10)
    arduino, ledning = g
    assert (arduino["bestilt"], arduino["antatt"], arduino["type"], arduino["utdelt"]) == (8.0, False, "utstyr", 10)
    assert (ledning["bestilt"], ledning["antatt"], ledning["type"]) == (20, True, "forbruk")


def test_forventet():
    l = {"lagerFor": 4, "bestilt": 6, "utdelt": 10}
    assert telling.forventet({**l, "type": "utstyr"}) == 10  # alt skal tilbake
    assert telling.forventet({**l, "type": "forbruk"}) == 0  # alt er brukt opp
    assert telling.forventet({"lagerFor": 50, "bestilt": 0, "utdelt": 20, "type": "forbruk"}) == 30


def test_lag_og_oppsummering():
    g = telling.grunnlag([rad("Arduino", "Elfa", 1, 4, 6, pris=250.0, varenr="A1"),
                          rad("Ledning", "Kjell", 2, 0, 20, varenr="L1"),
                          rad("Skrue", "", 1, 0, 10, varenr="S1")], ({}, set()), {}, sett=10)
    linjer = telling.lag(g, {"elfa|a1": {"type": "utstyr", "telt": 8}, "kjell|l1": {"type": "forbruk", "telt": 3}})
    assert [l["key"] for l in linjer] == ["elfa|a1", "kjell|l1"]  # skrue er ikke telt
    o = telling.oppsummering(linjer)
    # Arduino: forventet 10, telt 8 → svinn 2 × 250; ledning: forventet 0, telt 3 → avvik, ikke svinn
    assert o == {"telte": 2, "avvik": 2, "svinn": 1, "svinnKr": 500.0}
    with pytest.raises(ValueError):
        telling.lag(g, {"elfa|a1": {"type": "noe", "telt": 1}})
    with pytest.raises(ValueError):
        telling.lag(g, {"elfa|a1": {"telt": -1}})


# ---------- skriving til BOM ----------

def test_write_stock_endrer_bare_lager():
    path = lag_bom(rader=[("Arduino", "Elfa", "A1", 1, 4, 250.0), ("Ledning", "Kjell", "L1", 2, 0, 10.0)])
    backup_dir = config.BOM_DIR / "Sikkerhetskopi"
    res = bom.write_stock(path, {"elfa|a1": 9, "kjell|l1": 2.5, "finnes-ikke": 1}, backup_dir)
    assert res["skrevet"] == ["elfa|a1", "kjell|l1"] and res["hoppet"] == ["finnes-ikke"]
    assert (backup_dir / res["backup"]).exists()

    b = bom.read_bom(path)
    assert [(r["beskrivelse"], r["lager"], r["prSett"], r["pris"]) for r in b["rader"]] == [
        ("Arduino", 9, 1, 250.0), ("Ledning", 2.5, 2, 10.0)]
    ws = openpyxl.load_workbook(path).worksheets[0]
    assert str(ws["I14"].value).startswith("=")  # formlene står
    assert str(ws["K14"].value).startswith("=")
    # sikkerhetskopien har de gamle tallene
    assert bom.read_bom(backup_dir / res["backup"])["rader"][0]["lager"] == 4


def test_write_stock_hopper_over_doble_linjer():
    path = lag_bom(rader=[("Arduino", "Elfa", "A1", 1, 4, 250.0), ("Arduino igjen", "Elfa", "A1", 1, 1, 250.0)])
    res = bom.write_stock(path, {"elfa|a1": 9}, config.BOM_DIR / "Sikkerhetskopi")
    assert res["hoppet"] == ["elfa|a1"]
    assert [r["lager"] for r in bom.read_bom(path)["rader"]] == [4, 1]


def test_write_stock_nekter_naar_filen_er_aapen_i_excel():
    path = lag_bom(rader=[("Arduino", "Elfa", "A1", 1, 4, 250.0)])
    (path.parent / f"~${path.name}").write_text("lås")
    with pytest.raises(bom.BomError, match="åpen i Excel"):
        bom.write_stock(path, {"elfa|a1": 9}, config.BOM_DIR / "Sikkerhetskopi")
    assert bom.read_bom(path)["rader"][0]["lager"] == 4


# ---------- gjennom API-et ----------

def oppsett(client):
    g = client.post(G, json={"navn": "Test", "skolear": 2026, "start": 44, "sett": 10}).get_json()
    fil = lag_bom(rader=[("Arduino", "Elfa", "A1", 1, 4, 250.0), ("Ledning", "Kjell", "L1", 2, 0, 10.0)]).name
    client.patch(f"{G}/{g['id']}", json={"bom": fil})
    return g["id"], fil


def tell(client):
    return client.get("/api/state").get_json()["prosjekter"][0]["telling"]


def test_tellingsflyt(client):
    gid, fil = oppsett(client)
    client.post(f"{G}/{gid}/bestilling", json={"handling": "bestilt", "leverandor": "elfa"})
    t = tell(client)
    assert t["lagret"] is None and t["antattBestilt"]  # Kjell er ikke registrert
    assert [(l["key"], l["lagerFor"], l["bestilt"], l["utdelt"]) for l in t["grunnlag"]] == [
        ("elfa|a1", 4, 6, 10), ("kjell|l1", 0, 20, 20)]

    r = client.post(f"{G}/{gid}/telling", json={"linjer": [
        {"key": "elfa|a1", "type": "utstyr", "telt": 9}, {"key": "kjell|l1", "type": "forbruk", "telt": 4}]})
    assert r.get_json() == {"telte": 2, "avvik": 2, "svinn": 1, "svinnKr": 250.0}
    L = tell(client)["lagret"]
    assert (L["svinnKr"], L["skrevet"]) == (250.0, None)
    # de telte linjene må følge med til nettleseren, ikke bare antallet
    assert [l["key"] for l in L["linjer"]] == ["elfa|a1", "kjell|l1"] and L["telte"] == 2

    # typen huskes for prosjektet (neste år)
    import app as arshjul
    assert arshjul.store.linjetyper()[client.get("/api/state").get_json()["prosjekter"][0]["prosjektId"]] == {
        "elfa|a1": "utstyr", "kjell|l1": "forbruk"}

    # endre en lagret telling
    client.post(f"{G}/{gid}/telling", json={"linjer": [{"key": "elfa|a1", "type": "utstyr", "telt": 10},
                                                         {"key": "kjell|l1", "type": "forbruk", "telt": 4}]})
    assert tell(client)["lagret"]["svinn"] == 0

    r = client.post(f"{G}/{gid}/telling/skriv")
    assert r.status_code == 200 and r.get_json()["skrevet"] == ["elfa|a1", "kjell|l1"]
    assert [x["lager"] for x in bom.read_bom(config.BOM_DIR / fil)["rader"]] == [10, 4]
    L = tell(client)["lagret"]
    assert L["skrevet"] and L["backup"]

    # vår egen skriving vises ikke som «BOM-en er endret»
    assert client.get("/api/state").get_json()["prosjekter"][0]["bomStatus"]["endringer"] == []

    # ingenting nytt å skrive, og kan ikke slettes mens noe står i BOM-en
    assert client.post(f"{G}/{gid}/telling/skriv").status_code == 400
    assert client.delete(f"{G}/{gid}/telling").status_code == 400
    assert (L["iBom"], L["ventende"], len(L["logg"])) == (2, 0, 1)


def test_neste_aar_bruker_ny_lagerbeholdning(client):
    gid, fil = oppsett(client)
    client.post(f"{G}/{gid}/telling", json={"linjer": [{"key": "elfa|a1", "telt": 10}, {"key": "kjell|l1", "type": "forbruk", "telt": 4}]})
    client.post(f"{G}/{gid}/telling/skriv")
    client.post("/api/skolear/2026/kopier")
    neste = [p for p in client.get("/api/state").get_json()["prosjekter"] if p["skolear"] == 2027][0]
    # 10 sett: Arduino 10 − 10 på lager = 0; Ledning 20 − 4 = 16
    assert neste["bomStatus"]["bestill"] == 1
    g = {l["key"]: l for l in neste["telling"]["grunnlag"]}
    assert (g["elfa|a1"]["lagerFor"], g["kjell|l1"]["type"]) == (10, "forbruk")  # typen er husket


def test_slett_og_ugyldig_telling(client):
    gid, _ = oppsett(client)
    assert client.post(f"{G}/{gid}/telling", json={"linjer": []}).status_code == 400
    assert client.post(f"{G}/{gid}/telling", json={"linjer": [{"key": "elfa|a1", "telt": -2}]}).status_code == 400
    assert client.post(f"{G}/{gid}/telling/skriv").status_code == 400  # ikke lagret
    client.post(f"{G}/{gid}/telling", json={"linjer": [{"key": "elfa|a1", "telt": 3}]})
    assert client.delete(f"{G}/{gid}/telling").status_code == 200
    assert tell(client)["lagret"] is None


# ---------- rette, skrive på nytt og angre ----------

def test_write_stock_konflikt_og_ingen_kopi_uten_endring():
    path = lag_bom(rader=[("Arduino", "Elfa", "A1", 1, 4, 250.0), ("Ledning", "Kjell", "L1", 2, 0, 10.0)])
    kopier = config.BOM_DIR / "Sikkerhetskopi"
    res = bom.write_stock(path, {"elfa|a1": 9, "kjell|l1": 3}, kopier, forventet={"elfa|a1": 5, "kjell|l1": 0})
    assert res["skrevet"] == ["kjell|l1"]
    assert res["konflikt"] == [{"key": "elfa|a1", "forventet": 5.0, "faktisk": 4.0}]
    assert [r["lager"] for r in bom.read_bom(path)["rader"]] == [4, 3]
    # bare konflikter: ingenting skrives, ingen sikkerhetskopi
    antall = len(list(kopier.iterdir()))
    res = bom.write_stock(path, {"elfa|a1": 9}, kopier, forventet={"elfa|a1": 5})
    assert res["skrevet"] == [] and res["backup"] is None and len(list(kopier.iterdir())) == antall


def test_sikkerhetskopier_overskrives_aldri():
    path = lag_bom(rader=[("Arduino", "Elfa", "A1", 1, 4, 250.0)])
    kopier = config.BOM_DIR / "Sikkerhetskopi"
    navn = {bom.write_stock(path, {"elfa|a1": n}, kopier)["backup"] for n in (5, 6, 7)}  # samme sekund
    assert len(navn) == 3
    assert sorted(bom.read_bom(kopier / n)["rader"][0]["lager"] for n in navn) == [4, 5, 6]


def lager(fil):
    return {r["varenummer"]: r["lager"] for r in bom.read_bom(config.BOM_DIR / fil)["rader"]}


def test_rett_telling_og_skriv_bare_endringen(client):
    gid, fil = oppsett(client)
    client.post(f"{G}/{gid}/telling", json={"linjer": [{"key": "elfa|a1", "telt": 9}, {"key": "kjell|l1", "type": "forbruk", "telt": 4}]})
    client.post(f"{G}/{gid}/telling/skriv")
    assert lager(fil) == {"A1": 9, "L1": 4}

    # rett Arduino til 7: «før»-tallene fra BOM-en er de opprinnelige, ikke de nye
    client.post(f"{G}/{gid}/telling", json={"linjer": [{"key": "elfa|a1", "telt": 7}, {"key": "kjell|l1", "type": "forbruk", "telt": 4}]})
    L = tell(client)["lagret"]
    assert (L["ventende"], L["iBom"]) == (1, 2)
    assert {l["key"]: l["lagerFor"] for l in L["linjer"]} == {"elfa|a1": 4, "kjell|l1": 0}
    assert L["svinnKr"] == 750.0  # forventet 10, telt 7

    r = client.post(f"{G}/{gid}/telling/skriv").get_json()
    assert r["skrevet"] == ["elfa|a1"] and r["konflikt"] == []
    assert lager(fil) == {"A1": 7, "L1": 4}
    L = tell(client)["lagret"]
    assert (L["ventende"], [x["handling"] for x in L["logg"]]) == (0, ["skrevet", "skrevet"])


def test_endret_i_excel_overskrives_ikke(client):
    gid, fil = oppsett(client)
    client.post(f"{G}/{gid}/telling", json={"linjer": [{"key": "elfa|a1", "telt": 9}, {"key": "kjell|l1", "type": "forbruk", "telt": 4}]})
    client.post(f"{G}/{gid}/telling/skriv")
    # ingeniøren endrer Arduino i Excel etterpå
    bom.write_stock(config.BOM_DIR / fil, {"elfa|a1": 12}, config.BOM_DIR / "Sikkerhetskopi")
    client.post(f"{G}/{gid}/telling", json={"linjer": [{"key": "elfa|a1", "telt": 8}, {"key": "kjell|l1", "type": "forbruk", "telt": 5}]})
    r = client.post(f"{G}/{gid}/telling/skriv").get_json()
    assert r["skrevet"] == ["kjell|l1"]
    assert [(k["beskrivelse"], k["faktisk"], k["forventet"]) for k in r["konflikt"]] == [("Arduino", 12, 9)]
    assert lager(fil) == {"A1": 12, "L1": 5}
    assert tell(client)["lagret"]["ventende"] == 1  # Arduino venter fortsatt


def test_angre_setter_tilbake_og_aapner_for_sletting(client):
    gid, fil = oppsett(client)
    client.post(f"{G}/{gid}/telling", json={"linjer": [{"key": "elfa|a1", "telt": 9}, {"key": "kjell|l1", "type": "forbruk", "telt": 4}]})
    client.post(f"{G}/{gid}/telling/skriv")
    r = client.post(f"{G}/{gid}/telling/angre").get_json()
    assert sorted(r["skrevet"]) == ["elfa|a1", "kjell|l1"]
    assert lager(fil) == {"A1": 4, "L1": 0}  # som før tellingen
    L = tell(client)["lagret"]
    assert (L["iBom"], L["skrevet"], L["ventende"], L["logg"][-1]["handling"]) == (0, None, 2, "angret")
    assert client.post(f"{G}/{gid}/telling/angre").status_code == 400
    assert client.delete(f"{G}/{gid}/telling").status_code == 200


def test_migrering_4_fyller_inn_det_som_er_skrevet(tmp_path, monkeypatch):
    import json as _json

    import db as dbmod
    alle = dbmod.MIGRERINGER
    monkeypatch.setattr(dbmod, "MIGRERINGER", alle[:3])
    d = dbmod.Db(tmp_path / "x.db")
    d.init()
    with d.tx() as con:
        con.execute("INSERT INTO prosjekt (id, navn) VALUES ('p', 'P')")
        con.execute("INSERT INTO gjennomforing (id, prosjekt_id, skolear, start, slutt) VALUES (1, 'p', 2026, 40, 40)")
        con.execute("INSERT INTO telling (gjennomforing_id, dato, linjer, skrevet, backup) VALUES (1, '2026-10-06', ?, '2026-10-06', 'k.xlsx')",
                    (_json.dumps([{"key": "elfa|a1", "telt": 9}, {"key": "kjell|l1", "telt": 4.5}]),))
        con.execute("INSERT INTO gjennomforing (id, prosjekt_id, skolear, start, slutt) VALUES (2, 'p', 2027, 40, 40)")
        con.execute("INSERT INTO telling (gjennomforing_id, dato, linjer) VALUES (2, '2026-10-06', '[]')")
    monkeypatch.setattr(dbmod, "MIGRERINGER", alle[:4])  # til og med 4 (6 fjerner tellinger)
    d.init()
    t = d.tellinger()
    assert t["1"]["skrevet_verdier"] == {"elfa|a1": 9, "kjell|l1": 4.5}
    assert t["1"]["logg"][0]["backup"] == "k.xlsx"
    assert (t["2"]["skrevet_verdier"], t["2"]["logg"]) == ({}, [])


def test_migrering_6_ny_linjenokkel(tmp_path, monkeypatch):
    """Bestillinger får leverandøren i nøkkelen; tellinger og utstyr/forbruk (uten leverandør) fjernes."""
    import json as _json

    import db as dbmod
    alle = dbmod.MIGRERINGER
    monkeypatch.setattr(dbmod, "MIGRERINGER", alle[:5])
    d = dbmod.Db(tmp_path / "x.db")
    d.init()
    with d.tx() as con:
        con.execute("INSERT INTO prosjekt (id, navn) VALUES ('p', 'P')")
        con.execute("INSERT INTO gjennomforing (id, prosjekt_id, skolear, start, slutt) VALUES (1, 'p', 2026, 40, 40)")
        for lev, linjer, lagt in (("rs-online", [{"key": "1234", "antall": 2, "pris": 10.0, "beskrivelse": "RS"}], {"1234": 1}),
                                  ("farnell", [{"key": "1234", "antall": 3, "pris": 5.0, "beskrivelse": "F"}], {}),
                                  ("", [{"key": "skrue", "antall": 4, "pris": 1.0, "beskrivelse": "Skrue"}], {})):
            con.execute("""INSERT INTO bestilling (gjennomforing_id, leverandor, navn, status, bestilt_dato, linjer, paa_lager)
                           VALUES (1, ?, ?, 'bestilt', '2026-10-06', ?, ?)""", (lev, lev, _json.dumps(linjer), _json.dumps(lagt)))
        con.execute("INSERT INTO telling (gjennomforing_id, dato, linjer) VALUES (1, '2026-10-06', '[]')")
        con.execute("INSERT INTO linjetype (prosjekt_id, key, type) VALUES ('p', '1234', 'forbruk')")
    monkeypatch.setattr(dbmod, "MIGRERINGER", alle)
    d.init()
    b = d.bestillinger()["1"]
    assert [l["key"] for l in b["rs-online"]["linjer"]] == ["rs-online|1234"]
    assert b["rs-online"]["linjer"][0]["antall"] == 2 and b["rs-online"]["linjer"][0]["beskrivelse"] == "RS"
    assert b["rs-online"]["paa_lager"] == {"rs-online|1234": 1}
    assert [l["key"] for l in b["farnell"]["linjer"]] == ["farnell|1234"] and b["farnell"]["paa_lager"] == {}
    assert [l["key"] for l in b[""]["linjer"]] == ["|skrue"]
    assert d.tellinger() == {} and d.linjetyper() == {}


def test_endret_i_excel_vises_og_kan_hentes_inn_i_tellingen(client):
    """Ingeniørens test: tell 21, rett til 18 og skriv, sett 30 i Excel – tellingen skal vise det."""
    gid, fil = oppsett(client)
    client.post(f"{G}/{gid}/telling", json={"linjer": [{"key": "elfa|a1", "telt": 21}, {"key": "kjell|l1", "type": "forbruk", "telt": 4}]})
    client.post(f"{G}/{gid}/telling/skriv")
    client.post(f"{G}/{gid}/telling", json={"linjer": [{"key": "elfa|a1", "telt": 18}, {"key": "kjell|l1", "type": "forbruk", "telt": 4}]})
    client.post(f"{G}/{gid}/telling/skriv")
    assert tell(client)["lagret"]["endretIExcel"] == {}
    bom.write_stock(config.BOM_DIR / fil, {"elfa|a1": 30}, config.BOM_DIR / "Sikkerhetskopi")  # «Excel»

    L = tell(client)["lagret"]
    assert L["endretIExcel"] == {"elfa|a1": 30}
    assert {l["key"]: l["telt"] for l in L["linjer"]}["elfa|a1"] == 18  # tellingen endres ikke av seg selv

    r = client.post(f"{G}/{gid}/telling/hent").get_json()
    assert r["hentet"] == 1
    L = tell(client)["lagret"]
    assert {l["key"]: l["telt"] for l in L["linjer"]}["elfa|a1"] == 30
    assert (L["endretIExcel"], L["ventende"], L["logg"][-1]["handling"]) == ({}, 0, "hentet")
    assert lager(fil)["A1"] == 30  # BOM-en er ikke rørt
    assert client.post(f"{G}/{gid}/telling/hent").status_code == 400  # ingenting mer å hente

    # angre etterpå setter tilbake til tallet før tellingen, siden BOM-en har det årshjulet kjenner
    client.post(f"{G}/{gid}/telling/angre")
    assert lager(fil)["A1"] == 4
