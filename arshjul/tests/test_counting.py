import openpyxl
import pytest

import bom
import config
import counting
from helpers import make_bom

G = "/api/runs"


def row(desc, supplier, per_set, stock, to_order, price=10.0, part_no=""):
    return {"description": desc, "supplier": supplier, "part_number": part_no, "perSet": per_set, "in_stock": stock,
            "to_order": to_order, "price": price, "total": to_order * price}


# ---------- logikken ----------

def test_basis_uses_registered_order_or_assumes_bom():
    rows = [row("Arduino", "Elfa", 1, 4, 6, part_no="A1"), row("Ledning", "Kjell", 2, 0, 20, part_no="L1")]
    # Elfa is registered (and 8 were ordered), Kjell isn't
    g = counting.basis(rows, ({"elfa|a1": 8.0}, {"elfa"}), {"kjell|l1": "consumable"}, sets=10)
    arduino, ledning = g
    assert (arduino["ordered"], arduino["assumed"], arduino["type"], arduino["handed_out"]) == (8.0, False, "equipment", 10)
    assert (ledning["ordered"], ledning["assumed"], ledning["type"]) == (20, True, "consumable")


def test_expected():
    l = {"stockBefore": 4, "ordered": 6, "handed_out": 10}
    assert counting.expected({**l, "type": "equipment"}) == 10  # everything should come back
    assert counting.expected({**l, "type": "consumable"}) == 0  # everything is used up
    assert counting.expected({"stockBefore": 50, "ordered": 0, "handed_out": 20, "type": "consumable"}) == 30


def test_make_and_summary():
    g = counting.basis([row("Arduino", "Elfa", 1, 4, 6, price=250.0, part_no="A1"),
                          row("Ledning", "Kjell", 2, 0, 20, part_no="L1"),
                          row("Skrue", "", 1, 0, 10, part_no="S1")], ({}, set()), {}, sets=10)
    lines = counting.make(g, {"elfa|a1": {"type": "equipment", "counted": 8}, "kjell|l1": {"type": "consumable", "counted": 3}})
    assert [l["key"] for l in lines] == ["elfa|a1", "kjell|l1"]  # the screw hasn't been counted
    o = counting.summary(lines)
    # Arduino: expected 10, counted 8 → loss 2 × 250; wire: expected 0, counted 3 → deviation, not loss
    assert o == {"counted_lines": 2, "deviation": 2, "loss": 1, "lossCost": 500.0}
    with pytest.raises(ValueError):
        counting.make(g, {"elfa|a1": {"type": "noe", "counted": 1}})
    with pytest.raises(ValueError):
        counting.make(g, {"elfa|a1": {"counted": -1}})


# ---------- writing to the BOM ----------

def test_write_stock_changes_only_stock():
    path = make_bom(rows=[("Arduino", "Elfa", "A1", 1, 4, 250.0), ("Ledning", "Kjell", "L1", 2, 0, 10.0)])
    backup_dir = config.BOM_DIR / "Sikkerhetskopi"
    res = bom.write_stock(path, {"elfa|a1": 9, "kjell|l1": 2.5, "finnes-ikke": 1}, backup_dir)
    assert res["written"] == ["elfa|a1", "kjell|l1"] and res["skipped"] == ["finnes-ikke"]
    assert (backup_dir / res["backup"]).exists()

    b = bom.read_bom(path)
    assert [(r["description"], r["in_stock"], r["perSet"], r["price"]) for r in b["rows"]] == [
        ("Arduino", 9, 1, 250.0), ("Ledning", 2.5, 2, 10.0)]
    ws = openpyxl.load_workbook(path).worksheets[0]
    assert str(ws["I14"].value).startswith("=")  # the formulas are intact
    assert str(ws["K14"].value).startswith("=")
    # the backup has the old numbers
    assert bom.read_bom(backup_dir / res["backup"])["rows"][0]["in_stock"] == 4


def test_write_stock_skips_duplicate_lines():
    path = make_bom(rows=[("Arduino", "Elfa", "A1", 1, 4, 250.0), ("Arduino igjen", "Elfa", "A1", 1, 1, 250.0)])
    res = bom.write_stock(path, {"elfa|a1": 9}, config.BOM_DIR / "Sikkerhetskopi")
    assert res["skipped"] == ["elfa|a1"]
    assert [r["in_stock"] for r in bom.read_bom(path)["rows"]] == [4, 1]


def test_write_stock_refuses_when_the_file_is_open_in_excel():
    path = make_bom(rows=[("Arduino", "Elfa", "A1", 1, 4, 250.0)])
    (path.parent / f"~${path.name}").write_text("lås")
    with pytest.raises(bom.BomError, match="åpen i Excel"):
        bom.write_stock(path, {"elfa|a1": 9}, config.BOM_DIR / "Sikkerhetskopi")
    assert bom.read_bom(path)["rows"][0]["in_stock"] == 4


# ---------- through the API ----------

def setup(client):
    g = client.post(G, json={"name": "Test", "school_year": 2026, "start": 44, "students": 10}).get_json()
    file = make_bom(rows=[("Arduino", "Elfa", "A1", 1, 4, 250.0), ("Ledning", "Kjell", "L1", 2, 0, 10.0)]).name
    client.patch(f"{G}/{g['id']}", json={"bom": file})
    return g["id"], file


def cnt(client):
    return client.get("/api/state").get_json()["projects"][0]["count"]


def test_count_flow(client):
    gid, file = setup(client)
    client.post(f"{G}/{gid}/order", json={"action": "ordered", "supplier": "elfa"})
    t = cnt(client)
    assert t["saved"] is None and t["assumedOrdered"]  # Kjell isn't registered
    assert [(l["key"], l["stockBefore"], l["ordered"], l["handed_out"]) for l in t["basis"]] == [
        ("elfa|a1", 4, 6, 10), ("kjell|l1", 0, 20, 20)]

    r = client.post(f"{G}/{gid}/count", json={"lines": [
        {"key": "elfa|a1", "type": "equipment", "counted": 9}, {"key": "kjell|l1", "type": "consumable", "counted": 4}]})
    assert r.get_json() == {"counted_lines": 2, "deviation": 2, "loss": 1, "lossCost": 250.0}
    L = cnt(client)["saved"]
    assert (L["lossCost"], L["written"]) == (250.0, None)
    # the counted lines must go along to the browser, not just the number of them
    assert [l["key"] for l in L["lines"]] == ["elfa|a1", "kjell|l1"] and L["counted_lines"] == 2

    # the type is remembered for the project (next year)
    import app as arshjul
    assert arshjul.store.line_types()[client.get("/api/state").get_json()["projects"][0]["projectId"]] == {
        "elfa|a1": "equipment", "kjell|l1": "consumable"}

    # change a saved count
    client.post(f"{G}/{gid}/count", json={"lines": [{"key": "elfa|a1", "type": "equipment", "counted": 10},
                                                         {"key": "kjell|l1", "type": "consumable", "counted": 4}]})
    assert cnt(client)["saved"]["loss"] == 0

    r = client.post(f"{G}/{gid}/count/write")
    assert r.status_code == 200 and r.get_json()["written"] == ["elfa|a1", "kjell|l1"]
    assert [x["in_stock"] for x in bom.read_bom(config.BOM_DIR / file)["rows"]] == [10, 4]
    L = cnt(client)["saved"]
    assert L["written"] and L["backup"]

    # our own write doesn't show up as "the BOM has changed"
    assert client.get("/api/state").get_json()["projects"][0]["bomStatus"]["changes"] == []

    # nothing new to write, and it can't be deleted while something is in the BOM
    assert client.post(f"{G}/{gid}/count/write").status_code == 400
    assert client.delete(f"{G}/{gid}/count").status_code == 400
    assert (L["inBom"], L["pending"], len(L["log"])) == (2, 0, 1)


def test_next_year_uses_the_new_stock(client):
    gid, file = setup(client)
    client.post(f"{G}/{gid}/count", json={"lines": [{"key": "elfa|a1", "counted": 10}, {"key": "kjell|l1", "type": "consumable", "counted": 4}]})
    client.post(f"{G}/{gid}/count/write")
    client.post("/api/school-years/2026/copy")
    next_ = [p for p in client.get("/api/state").get_json()["projects"] if p["school_year"] == 2027][0]
    # 10 sets: Arduino 10 − 10 in stock = 0; wire 20 − 4 = 16
    assert next_["bomStatus"]["to_order"] == 1
    g = {l["key"]: l for l in next_["count"]["basis"]}
    assert (g["elfa|a1"]["stockBefore"], g["kjell|l1"]["type"]) == (10, "consumable")  # the type is remembered


def test_delete_and_invalid_count(client):
    gid, _ = setup(client)
    assert client.post(f"{G}/{gid}/count", json={"lines": []}).status_code == 400
    assert client.post(f"{G}/{gid}/count", json={"lines": [{"key": "elfa|a1", "counted": -2}]}).status_code == 400
    assert client.post(f"{G}/{gid}/count/write").status_code == 400  # not saved
    client.post(f"{G}/{gid}/count", json={"lines": [{"key": "elfa|a1", "counted": 3}]})
    assert client.delete(f"{G}/{gid}/count").status_code == 200
    assert cnt(client)["saved"] is None


# ---------- correcting, writing again and undoing ----------

def test_write_stock_conflict_and_no_backup_without_change():
    path = make_bom(rows=[("Arduino", "Elfa", "A1", 1, 4, 250.0), ("Ledning", "Kjell", "L1", 2, 0, 10.0)])
    copy_year = config.BOM_DIR / "Sikkerhetskopi"
    res = bom.write_stock(path, {"elfa|a1": 9, "kjell|l1": 3}, copy_year, expected={"elfa|a1": 5, "kjell|l1": 0})
    assert res["written"] == ["kjell|l1"]
    assert res["conflict"] == [{"key": "elfa|a1", "expected": 5.0, "actual": 4.0}]
    assert [r["in_stock"] for r in bom.read_bom(path)["rows"]] == [4, 3]
    # only conflicts: nothing is written, no backup
    quantity = len(list(copy_year.iterdir()))
    res = bom.write_stock(path, {"elfa|a1": 9}, copy_year, expected={"elfa|a1": 5})
    assert res["written"] == [] and res["backup"] is None and len(list(copy_year.iterdir())) == quantity


def test_backups_are_never_overwritten():
    path = make_bom(rows=[("Arduino", "Elfa", "A1", 1, 4, 250.0)])
    copy_year = config.BOM_DIR / "Sikkerhetskopi"
    name = {bom.write_stock(path, {"elfa|a1": n}, copy_year)["backup"] for n in (5, 6, 7)}  # samme sekund
    assert len(name) == 3
    assert sorted(bom.read_bom(copy_year / n)["rows"][0]["in_stock"] for n in name) == [4, 5, 6]


def stock(file):
    return {r["part_number"]: r["in_stock"] for r in bom.read_bom(config.BOM_DIR / file)["rows"]}


def test_correct_count_and_write_only_the_change(client):
    gid, file = setup(client)
    client.post(f"{G}/{gid}/count", json={"lines": [{"key": "elfa|a1", "counted": 9}, {"key": "kjell|l1", "type": "consumable", "counted": 4}]})
    client.post(f"{G}/{gid}/count/write")
    assert stock(file) == {"A1": 9, "L1": 4}

    # correct Arduino to 7: the "before" numbers from the BOM are the original ones, not the new ones
    client.post(f"{G}/{gid}/count", json={"lines": [{"key": "elfa|a1", "counted": 7}, {"key": "kjell|l1", "type": "consumable", "counted": 4}]})
    L = cnt(client)["saved"]
    assert (L["pending"], L["inBom"]) == (1, 2)
    assert {l["key"]: l["stockBefore"] for l in L["lines"]} == {"elfa|a1": 4, "kjell|l1": 0}
    assert L["lossCost"] == 750.0  # forventet 10, telt 7

    r = client.post(f"{G}/{gid}/count/write").get_json()
    assert r["written"] == ["elfa|a1"] and r["conflict"] == []
    assert stock(file) == {"A1": 7, "L1": 4}
    L = cnt(client)["saved"]
    assert (L["pending"], [x["action"] for x in L["log"]]) == (0, ["written", "written"])


def test_changed_in_excel_is_not_overwritten(client):
    gid, file = setup(client)
    client.post(f"{G}/{gid}/count", json={"lines": [{"key": "elfa|a1", "counted": 9}, {"key": "kjell|l1", "type": "consumable", "counted": 4}]})
    client.post(f"{G}/{gid}/count/write")
    # the engineer changes Arduino in Excel afterwards
    bom.write_stock(config.BOM_DIR / file, {"elfa|a1": 12}, config.BOM_DIR / "Sikkerhetskopi")
    client.post(f"{G}/{gid}/count", json={"lines": [{"key": "elfa|a1", "counted": 8}, {"key": "kjell|l1", "type": "consumable", "counted": 5}]})
    r = client.post(f"{G}/{gid}/count/write").get_json()
    assert r["written"] == ["kjell|l1"]
    assert [(k["description"], k["actual"], k["expected"]) for k in r["conflict"]] == [("Arduino", 12, 9)]
    assert stock(file) == {"A1": 12, "L1": 5}
    assert cnt(client)["saved"]["pending"] == 1  # Arduino venter fortsatt


def test_undo_restores_and_allows_deleting(client):
    gid, file = setup(client)
    client.post(f"{G}/{gid}/count", json={"lines": [{"key": "elfa|a1", "counted": 9}, {"key": "kjell|l1", "type": "consumable", "counted": 4}]})
    client.post(f"{G}/{gid}/count/write")
    r = client.post(f"{G}/{gid}/count/undo").get_json()
    assert sorted(r["written"]) == ["elfa|a1", "kjell|l1"]
    assert stock(file) == {"A1": 4, "L1": 0}  # as before the count
    L = cnt(client)["saved"]
    assert (L["inBom"], L["written"], L["pending"], L["log"][-1]["action"]) == (0, None, 2, "undone")
    assert client.post(f"{G}/{gid}/count/undo").status_code == 400
    assert client.delete(f"{G}/{gid}/count").status_code == 200


def test_migration_4_fills_in_what_was_written(tmp_path, monkeypatch):
    """Migration 4, checked against the schema it ran on (migration 6 later removes all counts)."""
    import json as _json

    import db as dbmod
    all_migrations = dbmod.MIGRATIONS
    monkeypatch.setattr(dbmod, "MIGRATIONS", all_migrations[:3])
    d = dbmod.Db(tmp_path / "x.db")
    d.init()
    with d.tx() as con:  # the old tables and the old (Norwegian) JSON keys, as they were then
        con.execute("INSERT INTO prosjekt (id, navn) VALUES ('p', 'P')")
        con.execute("INSERT INTO gjennomforing (id, prosjekt_id, skolear, start, slutt) VALUES (1, 'p', 2026, 40, 40)")
        con.execute("INSERT INTO telling (gjennomforing_id, dato, linjer, skrevet, backup) VALUES (1, '2026-10-06', ?, '2026-10-06', 'k.xlsx')",
                    (_json.dumps([{"key": "elfa|a1", "telt": 9}, {"key": "kjell|l1", "telt": 4.5}]),))
        con.execute("INSERT INTO gjennomforing (id, prosjekt_id, skolear, start, slutt) VALUES (2, 'p', 2027, 40, 40)")
        con.execute("INSERT INTO telling (gjennomforing_id, dato, linjer) VALUES (2, '2026-10-06', '[]')")
    monkeypatch.setattr(dbmod, "MIGRATIONS", all_migrations[:4])
    d.init()
    with d.tx() as con:
        rows = {r[0]: (_json.loads(r[1]), _json.loads(r[2]))
                for r in con.execute("SELECT gjennomforing_id, skrevet_verdier, logg FROM telling")}
    assert rows[1][0] == {"elfa|a1": 9, "kjell|l1": 4.5}
    assert rows[1][1][0]["backup"] == "k.xlsx"
    assert rows[2] == ({}, [])


def test_migration_6_new_line_key(tmp_path, monkeypatch):
    """Orders get the supplier in the key; counts and equipment/consumable (no supplier) are removed.
    Then migration 16 gives everything English names – checked through today's code."""
    import json as _json

    import db as dbmod
    all_migrations = dbmod.MIGRATIONS
    monkeypatch.setattr(dbmod, "MIGRATIONS", all_migrations[:5])
    d = dbmod.Db(tmp_path / "x.db")
    d.init()
    with d.tx() as con:  # old tables, old JSON keys
        con.execute("INSERT INTO prosjekt (id, navn) VALUES ('p', 'P')")
        con.execute("INSERT INTO gjennomforing (id, prosjekt_id, skolear, start, slutt) VALUES (1, 'p', 2026, 40, 40)")
        for supplier, lines, added in (("rs-online", [{"key": "1234", "antall": 2, "pris": 10.0, "beskrivelse": "RS"}], {"1234": 1}),
                                       ("farnell", [{"key": "1234", "antall": 3, "pris": 5.0, "beskrivelse": "F"}], {}),
                                       ("", [{"key": "skrue", "antall": 4, "pris": 1.0, "beskrivelse": "Skrue"}], {})):
            con.execute("""INSERT INTO bestilling (gjennomforing_id, leverandor, navn, status, bestilt_dato, linjer, paa_lager)
                           VALUES (1, ?, ?, 'bestilt', '2026-10-06', ?, ?)""", (supplier, supplier, _json.dumps(lines), _json.dumps(added)))
        con.execute("INSERT INTO telling (gjennomforing_id, dato, linjer) VALUES (1, '2026-10-06', '[]')")
        con.execute("INSERT INTO linjetype (prosjekt_id, key, type) VALUES ('p', '1234', 'forbruk')")
    monkeypatch.setattr(dbmod, "MIGRATIONS", all_migrations)
    d.init()
    b = d.all_orders()["1"]
    assert [l["key"] for l in b["rs-online"]["lines"]] == ["rs-online|1234"]
    assert b["rs-online"]["lines"][0]["quantity"] == 2 and b["rs-online"]["lines"][0]["description"] == "RS"
    assert b["rs-online"]["stocked"] == {"rs-online|1234": 1} and b["rs-online"]["status"] == "ordered"
    assert [l["key"] for l in b["farnell"]["lines"]] == ["farnell|1234"] and b["farnell"]["stocked"] == {}
    assert [l["key"] for l in b[""]["lines"]] == ["|skrue"]
    assert d.counts() == {} and d.line_types() == {}


def test_changed_in_excel_shows_and_can_be_fetched_into_the_count(client):
    """The engineer's test: count 21, correct to 18 and write, set 30 in Excel – the count should show it."""
    gid, file = setup(client)
    client.post(f"{G}/{gid}/count", json={"lines": [{"key": "elfa|a1", "counted": 21}, {"key": "kjell|l1", "type": "consumable", "counted": 4}]})
    client.post(f"{G}/{gid}/count/write")
    client.post(f"{G}/{gid}/count", json={"lines": [{"key": "elfa|a1", "counted": 18}, {"key": "kjell|l1", "type": "consumable", "counted": 4}]})
    client.post(f"{G}/{gid}/count/write")
    assert cnt(client)["saved"]["changedInExcel"] == {}
    bom.write_stock(config.BOM_DIR / file, {"elfa|a1": 30}, config.BOM_DIR / "Sikkerhetskopi")  # "Excel"

    L = cnt(client)["saved"]
    assert L["changedInExcel"] == {"elfa|a1": 30}
    assert {l["key"]: l["counted"] for l in L["lines"]}["elfa|a1"] == 18  # the count doesn't change by itself

    r = client.post(f"{G}/{gid}/count/fetch").get_json()
    assert r["fetched"] == 1
    L = cnt(client)["saved"]
    assert {l["key"]: l["counted"] for l in L["lines"]}["elfa|a1"] == 30
    assert (L["changedInExcel"], L["pending"], L["log"][-1]["action"]) == ({}, 0, "fetched")
    assert stock(file)["A1"] == 30  # the BOM hasn't been touched
    assert client.post(f"{G}/{gid}/count/fetch").status_code == 400  # nothing more to fetch

    # undoing afterwards sets it back to the number before the count, since the BOM holds what the app knows
    client.post(f"{G}/{gid}/count/undo")
    assert stock(file)["A1"] == 4
