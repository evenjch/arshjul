"""Reading and writing BOM files (Excel) made from the template.

The Excel file is the master copy of the components: this module reads it and creates new templates.
Two exceptions, both always with a backup first:
- the «På lager» column (H) after counting and receipts, when the user has confirmed it (write_stock)
- the set count (B10), which the app owns – otherwise Excel computes «Å bestille» wrong (write_sets)
"""
from __future__ import annotations

import datetime as dt
import json
import shutil
from pathlib import Path

import openpyxl

FIRST_ROW, LAST_ROW = 14, 73


class BomError(Exception):
    """An error that can be shown to the user as it is."""


def _num(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ".").strip())
    except ValueError:
        return 0.0


def _txt(v) -> str:
    return "" if v is None else str(v).strip()


def read_bom(path: Path) -> dict:
    """Read a BOM made from the template. Doesn't compute the quantity to order –
    that's done with the app's set count (see calc)."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except PermissionError as e:
        raise BomError(f"«{path.name}» er låst. Lukk den i Excel og prøv igjen.") from e
    except Exception as e:  # noqa: BLE001
        raise BomError(f"Kunne ikke lese «{path.name}». Er det en .xlsx-fil?") from e
    ws = wb.worksheets[0]
    if _txt(ws["C13"].value).lower() != "beskrivelse":  # the template's header in C13 (Excel data, Norwegian)
        raise BomError(f"«{path.name}» ser ikke ut som en BOM laget fra malen.")

    date = ws["F7"].value or ws["E7"].value
    if isinstance(date, (dt.datetime, dt.date)):
        date = date.strftime("%d.%m.%Y")
    vat = ws["D10"].value
    rows = []
    for r in range(FIRST_ROW, LAST_ROW + 1):
        desc = _txt(ws[f"C{r}"].value)
        if not desc:
            continue
        rows.append({
            "description": desc,
            "category": _txt(ws[f"D{r}"].value),
            "supplier": _txt(ws[f"E{r}"].value),
            "part_number": _txt(ws[f"F{r}"].value),
            "perSet": _num(ws[f"G{r}"].value),
            "in_stock": _num(ws[f"H{r}"].value),
            "price": _num(ws[f"J{r}"].value),
            "comment": _txt(ws[f"L{r}"].value),
        })
    return {
        "file": path.name,
        "project": _txt(ws["B7"].value),
        "responsible": _txt(ws["D7"].value),
        "date": _txt(date),
        "sets": int(_num(ws["B10"].value)) or None,
        "vat": 0.25 if vat in (None, "") else _num(vat),
        "rows": rows,
        "changed": dt.datetime.fromtimestamp(path.stat().st_mtime).strftime("%d.%m.%Y %H:%M"),
    }


def calc(bom: dict, sets: int) -> dict:
    """Quantity to order = per set × sets − in stock."""
    rows = []
    for r in bom["rows"]:
        to_order = max(0.0, r["perSet"] * sets - r["in_stock"])
        rows.append({**r, "to_order": to_order, "total": round(to_order * r["price"], 2)})
    excl = round(sum(r["total"] for r in rows), 2)
    return {
        "rows": rows,
        "lines": len(rows),
        "to_order": sum(1 for r in rows if r["to_order"] > 0),
        "excl": excl,
        "incl": round(excl * (1 + bom["vat"]), 2),
    }


# ---------- endringer siden sist ----------

def row_key(r: dict) -> str:
    """The identity of a BOM line: supplier + part number (or the description when there's no part number).

    Part numbers are only unique at one supplier – RS and Farnell can use the same number for different parts.
    """
    supplier = (r.get("supplier") or "").strip().lower()
    return f"{supplier}|{(r.get('part_number') or r.get('description') or '').strip().lower()}"


def _unique(rows: list[dict]) -> list[tuple[str, dict]]:
    """(key, row) where equal keys get «#2», «#3» … in order, so two identical lines
    in the same BOM aren't mixed up."""
    sets: dict[str, int] = {}
    out = []
    for r in rows:
        k = row_key(r)
        sets[k] = sets.get(k, 0) + 1
        out.append((k if sets[k] == 1 else f"{k}#{sets[k]}", r))
    return out


# the supplier is part of the key: changing supplier shows up as one line removed and a new one
FIELDS = [("description", "beskrivelse"), ("perSet", "per sett"),  # (key, label users see)
          ("in_stock", "på lager"), ("price", "pris"), ("comment", "kommentar")]


def diff(old: list[dict], new: list[dict]) -> list[dict]:
    old_k, new_k = _unique(old), _unique(new)
    a, b = dict(old_k), dict(new_k)
    out = []
    for k, r in new_k:
        if k not in a:
            out.append({"type": "new", "description": r["description"], "part_number": r["part_number"], "perSet": r["perSet"]})
    for k, r in old_k:
        if k not in b:
            out.append({"type": "removed", "description": r["description"], "part_number": r["part_number"]})
    for k, r in new_k:
        o = a.get(k)
        if not o:
            continue
        for k, label in FIELDS:
            x, y = o.get(k, ""), r.get(k, "")
            changed = abs(_num(x) - _num(y)) > 1e-9 if isinstance(y, (int, float)) else _txt(x) != _txt(y)
            if changed:
                out.append({"type": "changed", "description": r["description"], "field": label, "from": x, "to": y})
    return out


# ---------- updating stock after a count ----------

def write_stock(path: Path, values_: dict[str, float], backup_dir: Path,
                expected: dict[str, float] | None = None, add: bool = False) -> dict:
    """Write new «På lager» numbers (column H) in an existing BOM.

    values_:   {line key: new quantity}, or with add: the quantity to add
               (negative subtracts). Only column H is changed.
    expected:  {line key: what the cell should hold now}. If it holds something else (changed in
               Excel since), the line isn't written but reported as a conflict.
    add:       receipt. Adds to what the cell holds now – right even if the cell has been
               changed in Excel. If the result would be negative, it's a conflict.
    Returns: written/skipped/conflict, «created» = {key: the number now in the cell}, and backup.
    Lines that don't exist, or whose key is on several rows, are skipped.
    Takes a backup first, but only if something is actually going to be written.
    """
    expected = expected or {}
    wb = _open_for_writing(path)
    ws = wb.worksheets[0]
    rows: dict[str, list[int]] = {}
    for r in range(FIRST_ROW, LAST_ROW + 1):
        desc = _txt(ws[f"C{r}"].value)
        if desc:
            rows.setdefault(row_key({"description": desc, "part_number": _txt(ws[f"F{r}"].value),
                                      "supplier": _txt(ws[f"E{r}"].value)}), []).append(r)
    written, skipped, conflict, created = [], [], [], {}
    for key, quantity in values_.items():
        matches = rows.get(key, [])
        if len(matches) != 1:
            skipped.append(key)
            continue
        cell = ws[f"H{matches[0]}"]
        if key in expected and abs(_num(cell.value) - float(expected[key])) > 1e-9:
            conflict.append({"key": key, "expected": float(expected[key]), "actual": _num(cell.value)})
            continue
        new = _num(cell.value) + float(quantity) if add else float(quantity)
        if new < -1e-9:
            conflict.append({"key": key, "expected": -float(quantity), "actual": _num(cell.value)})
            continue
        cell.value = int(new) if float(new).is_integer() else new
        written.append(key)
        created[key] = new
    if not written:
        return {"written": [], "skipped": skipped, "conflict": conflict, "created": {}, "backup": None}
    backup = _save_with_backup(wb, path, backup_dir)
    return {"written": written, "skipped": skipped, "conflict": conflict, "created": created, "backup": backup}


# ---------- the set count (B10) ----------

_NONE = object()


def write_sets(path: Path, sets: int | None, backup_dir: Path, expected=_NONE) -> dict:
    """Write the app's set count to B10, which the BOM's «Å bestille» formulas use.

    expected: what B10 should hold now (the app's previous number). If it holds something else
              (changed by hand in Excel), nothing is written, and it's reported as a conflict.
    Returns: {"written": bool, "for": the number that was there, "file": the number in B10 now, "backup": name or None}.
    """
    wb = _open_for_writing(path)
    cell = wb.worksheets[0]["B10"]
    current = int(_num(cell.value)) or None
    if current == (sets or None):
        return {"written": False, "for": current, "file": current, "backup": None}
    if expected is not _NONE and current != (expected or None):
        return {"written": False, "for": current, "file": current, "backup": None, "conflict": True}
    cell.value = sets or None
    return {"written": True, "for": current, "file": sets or None, "backup": _save_with_backup(wb, path, backup_dir)}


# ---------- felles for skriving ----------

def _open_for_writing(path: Path):
    """Open a BOM with its formulas (not the values), and refuse if it's open in Excel."""
    if (path.parent / f"~${path.name}").exists():
        raise BomError(f"«{path.name}» er åpen i Excel. Lukk den og prøv igjen.")
    try:
        wb = openpyxl.load_workbook(path)
    except PermissionError as e:
        raise BomError(f"«{path.name}» er låst. Lukk den i Excel og prøv igjen.") from e
    if _txt(wb.worksheets[0]["C13"].value).lower() != "beskrivelse":  # the template's header in C13 (Excel data, Norwegian)
        raise BomError(f"«{path.name}» ser ikke ut som en BOM laget fra malen.")
    return wb


def _save_with_backup(wb, path: Path, backup_dir: Path) -> str:
    """Back up the file as it is, then save. Returns the name of the backup."""
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = f"{path.stem} {dt.datetime.now():%Y-%m-%d %H%M%S}"
    backup, n = backup_dir / f"{stamp}{path.suffix}", 2
    while backup.exists():  # never overwrite an earlier backup
        backup, n = backup_dir / f"{stamp} ({n}){path.suffix}", n + 1
    shutil.copy2(path, backup)
    wb.calculation.fullCalcOnLoad = True  # Excel recomputes «Å bestille» and «Total»
    try:
        wb.save(path)
    except PermissionError as e:
        raise BomError(f"«{path.name}» er låst. Lukk den i Excel og prøv igjen.") from e
    return backup.name


# ---------- nye filer ----------

def make_template(template: Path, target: Path, *, project: str, date: dt.date,
                  deadline: dt.date, sets: int | None) -> None:
    if target.exists():
        raise BomError(f"«{target.name}» finnes allerede i BOM-mappen.")
    wb = openpyxl.load_workbook(template)
    ws = wb.worksheets[0]
    ws["B7"] = project
    ws["F7"] = dt.datetime.combine(date, dt.time())
    ws["F10"] = dt.datetime.combine(deadline, dt.time())
    ws["B10"] = sets or None
    wb.save(target)


def make_order_list(template: Path, target: Path, *, bom: dict, calculated: dict, project: str,
                    sets: int, list_id: str) -> None:
    rows = []
    for i, r in enumerate(calculated["rows"], start=1):
        rows.append({"id": i, "no": i, "description": r["description"], "category": r["category"],
                     "supplier": r["supplier"], "part_number": r["part_number"],
                     "quantity": r["to_order"], "price": r["price"], "total": r["total"], "comment": r["comment"]})
    now = dt.datetime.now()
    data = {
        "meta": {"project": bom["project"] or project, "responsible": bom["responsible"], "date": bom["date"],
                 "sets": str(sets) if sets else ""},
        "cols": [{"key": "no", "label": "Nr"}, {"key": "description", "label": "Beskrivelse"},
                 {"key": "category", "label": "Kategori"}, {"key": "supplier", "label": "Leverandør"},
                 {"key": "part_number", "label": "Varenummer"}, {"key": "quantity", "label": "Antall", "kind": "qty"},
                 {"key": "price", "label": "Pris pr enhet", "kind": "money"},
                 {"key": "total", "label": "Total", "kind": "money"}, {"key": "comment", "label": "Kommentar"}],
        "rows": rows,
        "totals": {"excl": calculated["excl"], "incl": calculated["incl"]},
        "source": bom["file"],
        "generated": f"{now.day}.{now.month}.{now.year} {now:%H:%M}",
        "id": list_id,
    }
    js = json.dumps(data, ensure_ascii=False).replace("<", "\\u003c")
    html = template.read_text(encoding="utf-8").replace("__BOM_JSON__", js, 1)
    target.write_text(html, encoding="utf-8")
